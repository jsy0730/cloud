# Style kit mirroring the original deck (20 x 11.25 in, Arial, grey/beige palette)
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy
C=dict(bg='F4F5F4',ink='1F2224',ink2='3A3F42',mute='6B7276',line='B4B9BB',pale='E3E5E4',beige='E8D5A0',blue='7C8E96',
       red='B5523B',white='FFFFFF')
def rgb(h): return RGBColor.from_string(h)
FONT='Arial'
class Deck:
    def __init__(s,total,base=None,keep=1):
        if base:
            s.p=Presentation(base)
            ids=s.p.slides._sldIdLst
            for sid in list(ids)[keep:]:
                s.p.part.drop_rel(sid.rId); ids.remove(sid)
            s.n=keep
        else:
            s.p=Presentation(); s.p.slide_width=Inches(20); s.p.slide_height=Inches(11.25); s.n=0
        s.total=total
    def slide(s,bg=None):
        sl=s.p.slides.add_slide(s.p.slide_layouts[0 if len(s.p.slide_layouts)==1 else 6]); s.n+=1
        f=sl.background.fill; f.solid(); f.fore_color.rgb=rgb(bg or C['bg'])
        return sl
    def save(s,path): s.p.save(path)
def box(sl,x,y,w,h,fill=None,line=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE,dash=False):
    sh=sl.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    if dash:
        from pptx.enum.dml import MSO_LINE
        sh.line.dash_style=MSO_LINE.DASH
    sh.shadow.inherit=False
    return sh
def text(sl,x,y,w,h,runs,size=18,color='ink',bold=False,align='l',anchor='t',fill=None,line=None,inset=0.03,spacing=None,wrap=True):
    """runs: str | list of paragraphs; paragraph: str | list of (text, {opts})"""
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)) if not fill and not line else box(sl,x,y,w,h,fill,line)
    tf=tb.text_frame; tf.word_wrap=wrap
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,Inches(inset))
    tf.vertical_anchor={'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE,'b':MSO_ANCHOR.BOTTOM}[anchor]
    paras=runs if isinstance(runs,list) else [runs]
    for i,para in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
        if spacing: p.line_spacing=spacing
        segs=para if isinstance(para,list) else [(para,{})]
        for seg in segs:
            if isinstance(seg,str): seg=(seg,{})
            t,o=seg; r=p.add_run(); r.text=t
            f=r.font; f.name=FONT; f.size=Pt(o.get('size',size)); f.bold=o.get('bold',bold)
            f.color.rgb=rgb(C.get(o.get('color',color),o.get('color',color)))
            rPr=r._r.get_or_add_rPr(); ea=rPr.find(qn('a:ea'))
            if ea is None:
                ea=rPr.makeelement(qn('a:ea'),{}); rPr.append(ea)
            ea.set('typeface',FONT)
    return tb
def hline(sl,x,y,w,color='mute',wt=0.01):
    return box(sl,x,y,w,wt,fill=C[color] if color in C else color)
def vline(sl,x,y,h,color='mute',wt=0.01):
    return box(sl,x,y,wt,h,fill=C[color] if color in C else color)
def line(sl,x1,y1,x2,y2,color='ink',w=1.0,dash=False):
    c=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=rgb(C.get(color,color)); c.line.width=Pt(w)
    if dash:
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style=MSO_LINE.DASH
    return c
def pic(sl,path,x,y,w=None,h=None):
    return sl.shapes.add_picture(path,Inches(x),Inches(y),Inches(w) if w else None,Inches(h) if h else None)
def pic_crop(sl,path,x,y,w,h):
    """place image filling w x h box with centre crop"""
    from PIL import Image
    im=Image.open(path); iw,ih=im.size; tr=w/h; ir=iw/ih
    p=sl.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    if ir>tr: c=(1-tr/ir)/2; p.crop_left=c; p.crop_right=c
    else: c=(1-ir/tr)/2; p.crop_top=c; p.crop_bottom=c
    return p
SEC={1:('壹','사업 배경','事業槪要'),2:('貳','대지 분석','垈地分析'),3:('參','참고 사례','事例分析'),4:('肆','계획과 조닝','計畫槪要'),5:('伍','법규 검토','法規檢討')}
def header(sl,sec,title,x0=1.15,w=17.7,tag=None):
    m,k,h=SEC[sec]
    box(sl,x0,0.62,0.35,0.35,fill=C['mute'])
    text(sl,x0-0.04,0.6,0.44,0.4,m,size=15,bold=True,color='bg',align='c',anchor='m')
    text(sl,x0+0.47,0.64,1.4,0.35,[[(k,{'color':'ink2'}),('  '+h,{'color':'mute'})]],size=16.5,wrap=False)
    if tag: text(sl,x0+w-6,0.64,6,0.35,tag,size=13,color='mute',align='r')
    text(sl,x0,1.02,w,0.72,title,size=34,bold=True,color='ink',anchor='m')
    hline(sl,x0,1.84,w,'mute')
def footer(deck,sl,src=None,x0=1.15,w=17.7):
    if src: text(sl,x0,10.36,w,0.3,'출처 · '+src,size=10,color='mute')
    text(sl,x0,10.66,3,0.3,'%02d / %02d'%(deck.n,deck.total),size=11,color='ink2')
    text(sl,x0+w-6,10.66,6,0.3,'전포동 복합주차장 계획',size=11,color='ink2',align='r')
    # bottom progress bar
    seg=20/6
    for i in range(6): box(sl,i*seg+0.02,11.15,seg-0.04,0.1,fill=C['line'])
def rows(sl,x,y,w,items,kw=1.75,size=15,rh=0.42,valcolor='ink'):
    """label/value table with hairlines"""
    for i,(k,v) in enumerate(items):
        yy=y+i*rh
        hline(sl,x,yy,kw-0.1,'line'); hline(sl,x+kw,yy,w-kw,'line')
        text(sl,x,yy+0.05,kw,rh-0.05,k,size=size,color='ink2')
        text(sl,x+kw,yy+0.05,w-kw,rh-0.05,v,size=size,color=valcolor)
    hline(sl,x,y+len(items)*rh,w,'line')
    return y+len(items)*rh
def chip(sl,x,y,w,h,label,fill='ink',color='bg',size=13,bold=True):
    return text(sl,x,y,w,h,label,size=size,color=color,bold=bold,align='c',anchor='m',fill=C.get(fill,fill))
def table(sl,x,y,colw,data,rh=0.4,size=13,head_fill='ink',head_color='bg',zebra=None,first_bold=True,aligns=None,rhs=None):
    """simple grid table drawn with boxes; data[0]=header"""
    yy=y
    for r,row in enumerate(data):
        h=rhs[r] if rhs else rh
        xx=x
        for c,cell in enumerate(row):
            w=colw[c]
            if r==0:
                text(sl,xx,yy,w-0.03,h,cell,size=size,bold=True,color=head_color,fill=C.get(head_fill,head_fill),anchor='m',inset=0.08,align='c' if c else 'l')
            else:
                fill=None
                if zebra and r%2==0: fill=C[zebra]
                if isinstance(cell,tuple): cell,opt=cell
                else: opt={}
                text(sl,xx,yy,w-0.03,h,cell,size=opt.get('size',size),bold=opt.get('bold',first_bold and c==0),color=opt.get('color','ink'),
                     fill=C.get(opt.get('fill',fill),opt.get('fill',fill)) if (opt.get('fill') or fill) else None,anchor='m',inset=0.08,
                     align=opt.get('align',(aligns[c] if aligns else 'l')))
            xx+=w
        yy+=h
        hline(sl,x,yy-0.005,sum(colw)-0.03,'line')
    return yy

def freeform(sl,pts_in,fill=None,line=None,lw=0.75,dash=False):
    """pts_in: list of (x_in,y_in) inches"""
    E=914400
    fb=sl.shapes.build_freeform(int(pts_in[0][0]*E),int(pts_in[0][1]*E),scale=1.0)
    fb.add_line_segments([(int(x*E),int(y*E)) for x,y in pts_in[1:]],close=True)
    sh=fb.convert_to_shape()
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(C.get(fill,fill))
    else: sh.fill.background()
    if line:
        sh.line.color.rgb=rgb(C.get(line,line)); sh.line.width=Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style=MSO_LINE.DASH
    else: sh.line.fill.background()
    sh.shadow.inherit=False
    return sh
def num(sl,x,y,d,label,fill='ink',color='bg',size=12):
    return text(sl,x,y,d,d,label,size=size,bold=True,color=color,align='c',anchor='m',fill=C.get(fill,fill),inset=0)
